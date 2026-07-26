import os
from docx import Document
from reportlab.pdfgen import canvas
from fastapi.testclient import TestClient
from app.main import app

# 1. Create dummy files
os.makedirs("test_files", exist_ok=True)
os.makedirs("outputs", exist_ok=True)
os.makedirs("template", exist_ok=True)

if not os.path.exists("template/CQP_Template.docx"):
    Document().save("template/CQP_Template.docx")

# TMP
doc_tmp = Document()
doc_tmp.add_paragraph("Test Method Procedure")
doc_tmp.save("test_files/TMP_dummy.docx")

# ACL
doc_acl = Document()
doc_acl.add_paragraph("Duty Profile: Automotive Traction (conditioning rates: 0.5C, 1.0C)")
doc_acl.add_paragraph("Notes:")
doc_acl.save("test_files/ACL_dummy.docx")

# Datasheet
c = canvas.Canvas("test_files/DATASHEET_dummy.pdf")
c.drawString(100, 100, "Manufactured by: Cygnus Cell Technologies")
c.save()

# 2. Test Client
client = TestClient(app)

with open("test_files/TMP_dummy.docx", "rb") as f_tmp, \
     open("test_files/ACL_dummy.docx", "rb") as f_acl, \
     open("test_files/DATASHEET_dummy.pdf", "rb") as f_pdf:
     
    response = client.post(
        "/generate",
        files={
            "tmp_file": ("TMP_dummy.docx", f_tmp, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "acl_file": ("ACL_dummy.docx", f_acl, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "datasheet_file": ("DATASHEET_dummy.pdf", f_pdf, "application/pdf"),
        },
        data={"market": "EU / UN-38.3 US / DOT Global"}
    )

print(f"Status Code: {response.status_code}")
if response.status_code != 200:
    print(f"Error: {response.text}")
else:
    print("Success! Generated document size:", len(response.content))
