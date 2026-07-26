import pytest
from docx import Document
from app.generator.docx_parser import parse_acl, parse_tmp

def test_parse_tmp(tmp_path):
    docx_file = tmp_path / "tmp.docx"
    doc = Document()
    doc.add_paragraph("Test Method Procedure - CYG")
    doc.add_paragraph("NMC811 | Test")
    doc.save(str(docx_file))
    
    title, chem = parse_tmp(str(docx_file))
    assert "Test Method Procedure" in title
    assert "NMC811" in chem

def test_parse_acl(tmp_path):
    docx_file = tmp_path / "acl.docx"
    doc = Document()
    doc.add_paragraph("Acceptance Criteria & Limits")
    doc.add_paragraph("Duty Profile: Traction (conditioning rates: 1.0C)")
    doc.add_paragraph("Notes:")
    doc.add_paragraph("* Note 1")
    doc.save(str(docx_file))
    
    profiles, notes, title = parse_acl(str(docx_file))
    assert len(profiles) == 1
    assert profiles[0]['name'] == "Traction"
    assert "1.0C" in profiles[0]['rates']
    assert len(notes) == 1
