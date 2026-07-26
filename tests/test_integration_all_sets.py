import os
import pytest
import docx
from app.generator.core import generate_protocol

def test_integration_set_a(tmp_path):
    out_file = tmp_path / "SetA_test.docx"
    generate_protocol(
        tmp_path="test_files/inputs/setA/TMP_CYG-21700-50G.docx",
        acl_path="test_files/inputs/setA/ACL_CYG-21700-50G.docx",
        datasheet_path="test_files/inputs/setA/DATASHEET_CYG-21700-50G.pdf",
        market="EU / UN-38.3",
        output_path=str(out_file),
        template_path="test_files/template/CQP_Template.docx"
    )
    assert out_file.exists()
    doc = docx.Document(str(out_file))
    assert len(doc.tables) >= 6
    assert "CQP-CYG2170050G-00" in doc.tables[0].rows[0].cells[1].text
    assert "Cygnus Cell Technologies" in doc.tables[1].rows[0].cells[1].text

def test_integration_set_b(tmp_path):
    out_file = tmp_path / "SetB_test.docx"
    generate_protocol(
        tmp_path="test_files/inputs/setB/TMP_AUR-PR-340.docx",
        acl_path="test_files/inputs/setB/ACL_AUR-PR-340.docx",
        datasheet_path="test_files/inputs/setB/DATASHEET_AUR-PR-340.pdf",
        market="US / DOT",
        output_path=str(out_file),
        template_path="test_files/template/CQP_Template.docx"
    )
    assert out_file.exists()
    doc = docx.Document(str(out_file))
    assert len(doc.tables) >= 5
    assert "CQP-AURPR340-00" in doc.tables[0].rows[0].cells[1].text
    assert "Auriga Energy" in doc.tables[1].rows[0].cells[1].text

def test_integration_set_c(tmp_path):
    out_file = tmp_path / "SetC_test.docx"
    generate_protocol(
        tmp_path="test_files/inputs/setC/TMP_PLX-PCH-088.docx",
        acl_path="test_files/inputs/setC/ACL_PLX-PCH-088.docx",
        datasheet_path="test_files/inputs/setC/DATASHEET_PLX-PCH-088.pdf",
        market="Global",
        output_path=str(out_file),
        template_path="test_files/template/CQP_Template.docx"
    )
    assert out_file.exists()
    doc = docx.Document(str(out_file))
    assert len(doc.tables) >= 7
    assert "CQP-PLXPCH088-00" in doc.tables[0].rows[0].cells[1].text
    assert "Pollux Power" in doc.tables[1].rows[0].cells[1].text
