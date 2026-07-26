import docx
from app.generator.extractors.duty_profiles import extract_profiles_from_docx
from app.generator.extractors.test_parameters import extract_tests_from_table
from app.generator.extractors.footnotes import extract_footnotes

def parse_acl(acl_path: str):
    doc = docx.Document(acl_path)
    profiles = extract_profiles_from_docx(doc)
    footnotes = extract_footnotes(doc)
    
    # Map each ACL table to its corresponding duty profile
    for i, profile in enumerate(profiles):
        if i < len(doc.tables):
            profile['tests'] = extract_tests_from_table(doc.tables[i])
        else:
            profile['tests'] = []
            
    acl_doc_title = doc.paragraphs[0].text.strip() if doc.paragraphs else "Acceptance Criteria & Limits"
    return profiles, footnotes, acl_doc_title

def parse_tmp(tmp_path: str):
    doc = docx.Document(tmp_path)
    tmp_doc_title = doc.paragraphs[0].text.strip() if doc.paragraphs else "Test Method Procedure"
    
    # Extract cell chemistry from initial paragraphs
    chemistry = ""
    for p in doc.paragraphs[:5]:
        if "NMC" in p.text or "LFP" in p.text or "Graphite" in p.text:
            parts = p.text.split("|")
            if len(parts) > 2:
                chemistry = parts[2].strip()
                break
            chemistry = p.text.strip()
            
    return tmp_doc_title, chemistry
