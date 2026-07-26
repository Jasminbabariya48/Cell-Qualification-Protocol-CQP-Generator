import docx
from app.generator.pdf_parser import parse_datasheet
from app.generator.docx_parser import parse_acl, parse_tmp
from app.generator.template_engine import render_template
from app.generator.docx_builder import post_process_docx
from app.models.cell_data import CellQualificationData
from app.models.duty_profile import DutyProfile, TestParameter

def auto_correct_file_paths(tmp_path: str, acl_path: str):
    # Swap tmp and acl paths if user uploaded them in reverse order
    try:
        doc_a = docx.Document(tmp_path)
        doc_b = docx.Document(acl_path)
        
        a_has_tables = len(doc_a.tables) > 0
        b_has_tables = len(doc_b.tables) > 0
        
        a_title = doc_a.paragraphs[0].text if doc_a.paragraphs else ""
        b_title = doc_b.paragraphs[0].text if doc_b.paragraphs else ""
        
        if (a_has_tables and not b_has_tables) or ("Acceptance Criteria" in a_title or "Test Method Procedure" in b_title):
            return acl_path, tmp_path
    except Exception:
        pass
    return tmp_path, acl_path

def generate_protocol(tmp_path, acl_path, datasheet_path, market, output_path, template_path):
    # Ensure correct file ordering
    tmp_path, acl_path = auto_correct_file_paths(tmp_path, acl_path)
    
    # Parse source files
    metrics = parse_datasheet(datasheet_path)
    profiles_data, footnotes, acl_title = parse_acl(acl_path)
    tmp_title, chemistry_tmp = parse_tmp(tmp_path)
    
    # Build cell metadata and document ID
    cell_model = metrics.cell_model if metrics.cell_model else "CYG-21700-50G"
    chemistry = metrics.chemistry if metrics.chemistry else chemistry_tmp
    clean_model_id = cell_model.replace("-", "").replace(" ", "").upper()
    doc_number = f"CQP-{clean_model_id}-00"
    
    duty_profiles = []
    for pd in profiles_data:
        tests = [TestParameter(**t) for t in pd.get('tests', [])]
        duty_profiles.append(DutyProfile(name=pd['name'], rates=pd['rates'], tests=tests))
        
    duty_profiles_str = " & ".join([p.name for p in duty_profiles])
    
    cqp_data = CellQualificationData(
        doc_number=doc_number,
        market=market,
        cell_model=cell_model,
        chemistry=chemistry,
        tmp_doc_title=tmp_title,
        acl_doc_title=acl_title,
        duty_profiles_str=duty_profiles_str,
        metrics=metrics,
        duty_profiles=duty_profiles,
        footnotes=footnotes
    )
    
    data_dict = cqp_data.dict()
    for k, v in data_dict['metrics'].items():
        data_dict[k] = v
        
    data_dict['raw_duty_profiles'] = profiles_data
    data_dict['duty_profiles'] = duty_profiles_str
    
    # Preprocess XML template and post-process output Word file
    render_template(template_path, data_dict, output_path)
    post_process_docx(output_path, data_dict)
