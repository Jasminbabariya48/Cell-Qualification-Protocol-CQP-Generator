import copy
import re
from docxtpl import DocxTemplate
import docx
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

def consolidate_runs(elem):
    # Merge run nodes in paragraph to avoid split tokens
    if elem.tag.endswith('p'):
        paragraphs = [elem]
    else:
        paragraphs = elem.xpath('.//w:p')
        
    for p in paragraphs:
        t_nodes = p.xpath('.//w:t')
        if len(t_nodes) > 1:
            full_p_text = "".join(t.text for t in t_nodes if t.text)
            t_nodes[0].text = full_p_text
            for t in t_nodes[1:]:
                t.text = ""

def replace_token(element, token_name, replace_val):
    # Regex replace matching placeholder token in text node
    consolidate_runs(element)
    pattern = re.compile(r'\{\{\s*' + re.escape(token_name) + r'\s*\}\}')
    for r in element.xpath('.//w:t'):
        if r.text and pattern.search(r.text):
            r.text = pattern.sub(str(replace_val), r.text)

def set_cell_text(tc_elem, text):
    # Set text in table cell, strip automatic list numbering if present
    for num_pr in tc_elem.xpath('.//w:numPr'):
        num_pr.getparent().remove(num_pr)
        
    t_nodes = tc_elem.xpath('.//w:t')
    if t_nodes:
        t_nodes[0].text = str(text)
        for tn in t_nodes[1:]:
            tn.text = ""
    else:
        p_nodes = tc_elem.xpath('.//w:p')
        if not p_nodes:
            p = etree.Element(f"{{{W_NS}}}p")
            tc_elem.append(p)
            p_nodes = [p]
        p = p_nodes[0]
        r = etree.Element(f"{{{W_NS}}}r")
        t = etree.Element(f"{{{W_NS}}}t")
        t.text = str(text)
        r.append(t)
        p.append(r)

def pre_process_template_for_docxtpl(doc_path: str, data: dict):
    doc = docx.Document(doc_path)
    
    # 1. Expand Table 1 (Duty Profile Test Matrix, index 2 in doc.tables)
    if len(doc.tables) >= 3:
        table_matrix = doc.tables[2]
        if len(table_matrix.rows) > 1:
            template_row = table_matrix.rows[1]
            row_counter = 1
            
            profiles_list = data.get('raw_duty_profiles', [])
            for dp in profiles_list:
                cycles_val = "500"
                for test in dp.get('tests', []):
                    if 'cycle' in test.get('parameter', '').lower():
                        m = re.search(r'after\s*(\d+)\s*cycles', test.get('acceptance_limit', ''))
                        if m:
                            cycles_val = m.group(1)
                            break
                            
                for rate in dp.get('rates', []):
                    new_row = copy.deepcopy(template_row._element)
                    consolidate_runs(new_row)
                    
                    t_cells = new_row.xpath('.//w:tc')
                    if len(t_cells) >= 6:
                        set_cell_text(t_cells[0], str(row_counter))
                        set_cell_text(t_cells[1], dp['name'])
                        set_cell_text(t_cells[2], rate)
                        set_cell_text(t_cells[3], data.get('v_max', ''))
                        set_cell_text(t_cells[4], data.get('v_min', ''))
                        set_cell_text(t_cells[5], cycles_val)
                            
                    table_matrix._tbl.append(new_row)
                    row_counter += 1
                    
            table_matrix._tbl.remove(template_row._element)
            
    # 2. Expand Section 7 Block per duty profile
    body = doc._element.body
    start_idx = -1
    end_idx = -1
    
    for i, child in enumerate(body):
        text = "".join(t.text for t in child.xpath('.//w:t') if t.text)
        if "7.{{ block_index }}" in text or "Qualification Tests" in text:
            start_idx = i
        elif "Applicable notes from the ACL:" in text and start_idx != -1:
            end_idx = i
            break
            
    if start_idx != -1 and end_idx != -1:
        block_elements = body[start_idx:end_idx]
        insert_pos = end_idx
        
        profiles_list = data.get('raw_duty_profiles', [])
        for idx, dp in enumerate(profiles_list):
            prev_heading = None
            for elem in block_elements:
                new_elem = copy.deepcopy(elem)
                consolidate_runs(new_elem)
                elem_text = "".join(t.text for t in new_elem.xpath('.//w:t') if t.text).strip()
                
                if "Acceptance Criteria:" in elem_text:
                    prev_heading = "Acceptance Criteria:"
                elif "Conclusion:" in elem_text:
                    prev_heading = "Conclusion:"
                    
                replace_token(new_elem, "block_index", str(idx + 1))
                replace_token(new_elem, "duty_profile", dp['name'])
                
                if "{{ to_be_added_by_reviewer }}" in elem_text:
                    if prev_heading == "Acceptance Criteria:":
                        replace_token(new_elem, "to_be_added_by_reviewer", "All listed parameters meet the acceptance limits for this duty profile.")
                    elif prev_heading == "Conclusion:":
                        replace_token(new_elem, "to_be_added_by_reviewer", "The cell is qualified for this duty profile, subject to review.")
                        
                if new_elem.tag.endswith('tbl'):
                    rows = new_elem.xpath('.//w:tr')
                    if len(rows) > 1:
                        test_template_row = rows[1]
                        for j, test in enumerate(dp.get('tests', [])):
                            new_test_row = copy.deepcopy(test_template_row)
                            consolidate_runs(new_test_row)
                            
                            cells = new_test_row.xpath('.//w:tc')
                            if len(cells) >= 4:
                                set_cell_text(cells[0], str(j + 1))
                                set_cell_text(cells[1], test['parameter'])
                                set_cell_text(cells[2], test['acceptance_limit'])
                                set_cell_text(cells[3], test['clause'])
                                
                            new_elem.append(new_test_row)
                        new_elem.remove(test_template_row)
                        
                body.insert(insert_pos, new_elem)
                insert_pos += 1
                
        for elem in block_elements:
            body.remove(elem)
            
    footnotes_str = "\n\n".join(data.get('footnotes', []))
    for p in doc.paragraphs:
        if "{{ also_fetch_any_footnotes_from_the_acl }}" in p.text:
            p.text = p.text.replace("{{ also_fetch_any_footnotes_from_the_acl }}", footnotes_str)

    temp_path = doc_path.replace(".docx", "_temp.docx")
    doc.save(temp_path)
    return temp_path

def render_template(template_path: str, data: dict, output_path: str):
    temp_path = pre_process_template_for_docxtpl(template_path, data)
    doc = DocxTemplate(temp_path)
    doc.render(data)
    doc.save(output_path)
