from lxml import etree
import docx

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_MAP = {'w': W_NS}

def apply_vertical_merge(cell, merge_type="continue"):
    # Apply vertical cell merging in docx table using lxml XML elements
    tcPr = cell._tc.get_or_add_tcPr()
    vMerge = etree.Element(f"{{{W_NS}}}vMerge")
    vMerge.set(f"{{{W_NS}}}val", merge_type)
    tcPr.append(vMerge)

def protect_document(docx_path: str):
    # Enforce read-only protection while leaving reviewer text zones editable
    doc = docx.Document(docx_path)
    
    settings = doc.settings.element
    document_protection = etree.Element(f"{{{W_NS}}}documentProtection")
    document_protection.set(f"{{{W_NS}}}edit", "readOnly")
    document_protection.set(f"{{{W_NS}}}enforcement", "1")
    settings.append(document_protection)
    
    perm_id = 0
    for p in doc.paragraphs:
        if "All listed parameters meet" in p.text or "subject to review" in p.text:
            perm_id += 1
            permStart = etree.Element(f"{{{W_NS}}}permStart")
            permStart.set(f"{{{W_NS}}}id", str(perm_id))
            permStart.set(f"{{{W_NS}}}edGrp", "everyone")
            
            p._element.insert(0, permStart)
            
            permEnd = etree.Element(f"{{{W_NS}}}permEnd")
            permEnd.set(f"{{{W_NS}}}id", str(perm_id))
            p._element.append(permEnd)

    doc.save(docx_path)
    
def post_process_docx(docx_path: str, data: dict):
    doc = docx.Document(docx_path)
    
    # Vertically merge duty profile names in Table 1 matrix across rate rows
    if len(doc.tables) >= 3:
        table_matrix = doc.tables[2]
        current_profile = None
        
        for row_idx, row in enumerate(table_matrix.rows):
            if row_idx == 0:
                continue
            
            profile_name = row.cells[1].text.strip()
            
            if profile_name != current_profile:
                current_profile = profile_name
                apply_vertical_merge(row.cells[1], "restart")
            else:
                apply_vertical_merge(row.cells[1], "continue")
                row.cells[1].text = ""

    doc.save(docx_path)
    protect_document(docx_path)
